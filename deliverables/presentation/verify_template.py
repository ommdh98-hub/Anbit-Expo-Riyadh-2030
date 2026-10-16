from pathlib import Path
from zipfile import ZipFile
from lxml import etree as E
from pptx import Presentation
import hashlib,json,re
root=Path(__file__).resolve().parent
a=root/'template.pptx';b=root/'ANBIT-Expo-Riyadh-2030.pptx';ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
def sig(e):
 if e is None:return None
 return [e.tag,sorted(e.attrib.items()),e.text or '',[sig(x) for x in e]]
def walk(ss):
 for s in ss:
  yield s
  if s.shape_type==6:yield from walk(s.shapes)
pa=Presentation(a);pb=Presentation(b);assert len(pa.slides)==len(pb.slides)==14
report={'slide_count':14,'protected_slide_xml_equal':{},'shapes_unchanged':{},'run_properties_and_paragraph_properties':{},'untouched_parts_equal':True,'screenshot_status':'pending: requested directory absent','forbidden_text_found':[]}
alltext=[]
with ZipFile(a) as za,ZipFile(b) as zb:
 assert not zb.testzip()
 changed=[]
 for n in za.namelist():
  if za.read(n)!=zb.read(n):changed.append(n)
 allowed=['docProps/core.xml']+[f'ppt/slides/slide{i}.xml' for i in range(4,14)]
 assert all(n in allowed for n in changed),changed
 report['changed_parts']=changed
 for i in [1,2,3,14]:
  eq=za.read(f'ppt/slides/slide{i}.xml')==zb.read(f'ppt/slides/slide{i}.xml');assert eq;report['protected_slide_xml_equal'][str(i)]=eq
 for i,(sa,sb) in enumerate(zip(pa.slides,pb.slides),1):
  aa=list(walk(sa.shapes));bb=list(walk(sb.shapes));assert len(aa)==len(bb)
  for xa,xb in zip(aa,bb):
   assert (xa.shape_id,xa.shape_type,xa.left,xa.top,xa.width,xa.height)==(xb.shape_id,xb.shape_type,xb.left,xb.top,xb.width,xb.height),(i,xa.shape_id)
   ea=xa._element;eb=xb._element
   for tag in ['p:spPr','p:grpSpPr','p:style','p:blipFill']:
    assert sig(ea.find(tag,ns))==sig(eb.find(tag,ns)),(i,xa.shape_id,tag)
   if not xa.has_text_frame:continue
   assert sig(xa.text_frame._txBody.find('a:bodyPr',ns))==sig(xb.text_frame._txBody.find('a:bodyPr',ns))
   aps=xa.text_frame.paragraphs;bps=xb.text_frame.paragraphs;assert len(bps)<=len(aps)
   for j,(ap,bp) in enumerate(zip(aps,bps)):
    assert sig(ap._p.find('a:pPr',ns))==sig(bp._p.find('a:pPr',ns)),(i,xa.shape_id,j,'pPr')
    for k,br in enumerate(bp.runs):
     ar=ap.runs[k] if k<len(ap.runs) else None
     if ar is None:
      if i==4 and xa.shape_id==143:ar=next(x for x in aa if x.shape_id==144).text_frame.paragraphs[0].runs[0]
      elif i==8 and xa.shape_id==216:ar=aps[1].runs[0]
      else:raise AssertionError((i,xa.shape_id,j,k,'new unapproved run'))
     assert sig(ar._r.rPr)==sig(br._r.rPr),(i,xa.shape_id,j,k,'rPr')
   alltext.append(xb.text)
  report['shapes_unchanged'][str(i)]={'count_including_group_children':len(aa),'pass':True}
  report['run_properties_and_paragraph_properties'][str(i)]=True
 # All slide notes and file metadata are inspected, including empty notes.
 for n in zb.namelist():
  if n.endswith('.xml') and (n.startswith('ppt/notesSlides/') or n.startswith('docProps/')):
   x=E.fromstring(zb.read(n));alltext.extend(x.xpath('//text()'))
 text='\n'.join(alltext)
 banned=['يعرض هذا القسم','Stay ahead','Seasonal','quarter','عنوان','تذكر','Monitor changes','Identify times','ما هي','(على سبيل المثال','Codex','Claude','ChatGPT','Grow_Right','Grow Right']
 for t in banned:
  if t.casefold() in text.casefold():report['forbidden_text_found'].append(t)
 if '0%' in text:report['forbidden_text_found'].append('0%')
 if '!' in text:report['forbidden_text_found'].append('!')
 if re.search('[\U0001F300-\U0001FAFF\u2600-\u27BF]',text):report['forbidden_text_found'].append('emoji')
 assert not report['forbidden_text_found'],report['forbidden_text_found']
 core=E.fromstring(zb.read('docProps/core.xml'))
 assert core.find('{http://purl.org/dc/elements/1.1/}creator').text=='فريق أنبت'
report['template_sha256']=hashlib.sha256(a.read_bytes()).hexdigest()
report['output_sha256']=hashlib.sha256(b.read_bytes()).hexdigest()
report['status']='text and preservation checks passed; screenshot requirement pending'
(root/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
